"""STAGE 0a — extract seed tuples from the usage-source taxonomy workbook.

Deterministic, no LLM (spec §5 Stage 0a as amended by A-1 / D-0008):
reads seeds/usage_sources/llm_everyday_use_cases_taxonomy.xlsx and emits
seeds/seed_tuples.parquet — one row per (domain × user_type × use_case)
candidate tuple, unweighted, with per-row source provenance joined from the
workbook's Evidence map sheet.

Seed IDs are assigned in workbook row order, so re-running on the same
workbook (verified by sha256) is byte-stable.

Usage: python src/stage0_seeds.py
"""

from __future__ import annotations

import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pandas as pd

REPO = Path(__file__).resolve().parent.parent
WORKBOOK = REPO / "seeds/usage_sources/llm_everyday_use_cases_taxonomy.xlsx"
OUT = REPO / "seeds/seed_tuples.parquet"
SUMMARY = REPO / "seeds/seed_tuples_summary.json"

# Pin the exact workbook this stage was ratified against (D-0006).
EXPECTED_SHA256 = "b9fd845925ddd9ffb41276eb4fbedfb58d9badb76e22c89eac168c0446445cf4"

MIN_DISTINCT_TUPLES = 1200  # spec §5 Stage 0a floor: quartets_target / 8


def main() -> None:
    digest = hashlib.sha256(WORKBOOK.read_bytes()).hexdigest()
    if digest != EXPECTED_SHA256:
        sys.exit(
            f"Workbook sha256 mismatch:\n  got      {digest}\n  expected {EXPECTED_SHA256}\n"
            "The source file changed; re-ratify via a DECISIONS.md entry before rerunning."
        )

    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)

    use_rows = [
        r for r in wb["Use cases"].iter_rows(min_row=2, values_only=True) if any(r)
    ]
    evid = {
        int(r[0]): {"evidence_basis": str(r[3]).strip(), "source_key": str(r[4]).strip()}
        for r in wb["Evidence map"].iter_rows(min_row=2, values_only=True)
        if r[0] is not None
    }

    records = []
    for i, (domain, user, use_case) in enumerate(use_rows):
        workbook_row = i + 2  # 1-indexed sheet row, matching the Evidence map key
        e = evid.get(workbook_row)
        if e is None:
            sys.exit(f"Use-cases row {workbook_row} has no Evidence map entry — workbook integrity broken.")
        records.append(
            {
                "seed_tuple_id": f"seed-{i:06d}",
                "domain": str(domain).strip(),
                "user_type": str(user).strip(),
                "use_case": str(use_case).strip(),
                "workbook_row": workbook_row,
                "evidence_basis": e["evidence_basis"],
                "source_key": e["source_key"],
            }
        )

    df = pd.DataFrame(records)

    n_distinct = df.drop_duplicates(["domain", "user_type", "use_case"]).shape[0]
    if n_distinct != len(df):
        sys.exit(f"{len(df) - n_distinct} duplicate tuples — workbook integrity broken.")
    if n_distinct < MIN_DISTINCT_TUPLES:
        sys.exit(f"Only {n_distinct} tuples < floor {MIN_DISTINCT_TUPLES}.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    df.to_parquet(OUT, index=False)

    summary = {
        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "workbook_sha256": digest,
        "n_tuples": len(df),
        "n_domains": df["domain"].nunique(),
        "n_user_types": df["user_type"].nunique(),
        "n_domain_user_pairs": df.groupby(["domain", "user_type"]).ngroups,
        "tuples_per_domain": df["domain"].value_counts().to_dict(),
        "source_key_counts": df["source_key"].value_counts().to_dict(),
    }
    SUMMARY.write_text(json.dumps(summary, indent=2) + "\n")

    print(f"Wrote {OUT.relative_to(REPO)}: {len(df)} tuples, "
          f"{summary['n_domains']} domains, {summary['n_domain_user_pairs']} domain-user pairs.")
    print(f"Summary → {SUMMARY.relative_to(REPO)}")


if __name__ == "__main__":
    main()
